"""Minimal parsers: enough to establish real observations, not concept mapping."""

from __future__ import annotations

import csv
import io
import json
import math
import xml.etree.ElementTree as ET
from typing import Any


def numeric(value: Any) -> bool:
    try:
        return value is not None and math.isfinite(float(value))
    except (ValueError, TypeError):
        return False


def parse_csv(body: bytes) -> dict[str, Any]:
    text = body.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    columns = reader.fieldnames or []
    time_key = next(
        (x for x in ("TIME_PERIOD", "DATE", "observation_date", "date") if x in columns), None
    )
    value_key = next((x for x in ("OBS_VALUE", "value", "VALUE") if x in columns), None)
    if value_key is None and time_key is not None:
        remaining = [x for x in columns if x != time_key]
        if len(remaining) == 1:
            value_key = remaining[0]
    if time_key is None or value_key is None:
        raise ValueError("No recognised observation date/value columns")
    rows = [row for row in reader if row.get(time_key) and numeric(row.get(value_key))]
    if not rows:
        raise ValueError("CSV contains no finite observations")
    first = rows[0]
    metadata = {k: v for k, v in first.items() if k not in (time_key, value_key) and k is not None}
    return {
        "format": "csv",
        "observation_count": len(rows),
        "columns": columns,
        "first_period": min(row[time_key] for row in rows),
        "last_period": max(row[time_key] for row in rows),
        "example_observation": {"period": first[time_key], "value": float(first[value_key])},
        "series_metadata": metadata,
    }


def parse_sdmx_xml(body: bytes) -> dict[str, Any]:
    root = ET.fromstring(body)
    observations: list[dict[str, Any]] = []
    metadata: dict[str, str] = {}
    for element in root.iter():
        tag = element.tag.split("}")[-1]
        if tag == "Series" and not metadata:
            metadata.update(element.attrib)
            for child in element:
                if child.tag.split("}")[-1] in ("SeriesKey", "Attributes"):
                    metadata.update(
                        {
                            n.attrib["id"]: n.attrib["value"]
                            for n in child
                            if "id" in n.attrib and "value" in n.attrib
                        }
                    )
        if tag != "Obs":
            continue
        attrs = dict(element.attrib)
        for child in element:
            child_tag = child.tag.split("}")[-1]
            if child_tag in ("ObsDimension", "Time"):
                attrs["TIME_PERIOD"] = child.attrib.get("value", child.text or "")
            elif child_tag == "ObsValue":
                attrs["OBS_VALUE"] = child.attrib.get("value", "")
        if attrs.get("TIME_PERIOD") and numeric(attrs.get("OBS_VALUE")):
            observations.append(
                {"period": attrs["TIME_PERIOD"], "value": float(attrs["OBS_VALUE"])}
            )
    if not observations:
        raise ValueError("SDMX XML contains no finite observations")
    return {
        "format": "sdmx-xml",
        "observation_count": len(observations),
        "first_period": min(x["period"] for x in observations),
        "last_period": max(x["period"] for x in observations),
        "example_observation": observations[0],
        "series_metadata": metadata,
    }


def parse_sdmx(body: bytes) -> dict[str, Any]:
    return parse_sdmx_xml(body) if body.lstrip().startswith(b"<") else parse_csv(body)


def parse_worldbank(body: bytes) -> dict[str, Any]:
    payload = json.loads(body)
    if not isinstance(payload, list) or len(payload) != 2 or not isinstance(payload[1], list):
        raise ValueError("WDI response has no observation list")
    rows = [row for row in payload[1] if numeric(row.get("value"))]
    if not rows:
        raise ValueError("WDI response contains no finite observations")
    first = rows[0]
    return {
        "format": "wdi-json",
        "observation_count": len(rows),
        "first_period": min(row["date"] for row in rows),
        "last_period": max(row["date"] for row in rows),
        "example_observation": {"period": first["date"], "value": float(first["value"])},
        "series_metadata": {
            "indicator": first["indicator"],
            "country": first["country"],
            "unit": first.get("unit"),
            "sourceid": payload[0].get("sourceid"),
            "lastupdated": payload[0].get("lastupdated"),
        },
    }


def parse_eurostat(body: bytes) -> dict[str, Any]:
    payload = json.loads(body)
    values = payload.get("value", {})
    items = values.items() if isinstance(values, dict) else enumerate(values)
    observations = [(str(k), float(v)) for k, v in items if numeric(v)]
    if not observations or "dimension" not in payload:
        raise ValueError("JSON-stat response contains no finite observations or dimensions")
    time_labels = payload["dimension"]["time"]["category"]["label"]
    return {
        "format": "json-stat",
        "observation_count": len(observations),
        "first_period": min(time_labels),
        "last_period": max(time_labels),
        "example_observation": {"flat_index": observations[0][0], "value": observations[0][1]},
        "series_metadata": {
            "label": payload.get("label"),
            "updated": payload.get("updated"),
            "dimensions": {
                key: value.get("category")
                for key, value in payload["dimension"].items()
                if key != "time"
            },
        },
    }


def parse_pink_sheet(body: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = workbook["Monthly Prices"]
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next(i for i, row in enumerate(rows[:20]) if "Crude oil, Brent" in row)
        column = rows[header_index].index("Crude oil, Brent")
        observations: list[dict[str, Any]] = [
            {"period": str(row[0]), "value": float(str(row[column]))}
            for row in rows[header_index + 1 :]
            if row
            and isinstance(row[0], str)
            and "M" in row[0]
            and len(row) > column
            and numeric(row[column])
        ]
        if not observations:
            raise ValueError("Pink Sheet workbook contains no Brent monthly observations")
        return {
            "format": "xlsx",
            "observation_count": len(observations),
            "first_period": min(x["period"] for x in observations),
            "last_period": max(x["period"] for x in observations),
            "example_observation": observations[-1],
            "series_metadata": {
                "sheet": sheet.title,
                "column": "Crude oil, Brent",
                "header_row": header_index + 1,
            },
        }
    finally:
        workbook.close()


def parse_response(parser: str, body: bytes) -> dict[str, Any]:
    parsers = {
        "sdmx": parse_sdmx,
        "csv": parse_csv,
        "wdi": parse_worldbank,
        "jsonstat": parse_eurostat,
        "pink_sheet": parse_pink_sheet,
    }
    if parser not in parsers:
        raise ValueError(f"Unknown parser: {parser}")
    return parsers[parser](body)
