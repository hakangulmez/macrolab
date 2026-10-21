"""World Bank profiles and Pink Sheet downloads, with auditable live-link resolution."""

from __future__ import annotations

import io
import json
import math
import re
from datetime import UTC, date, datetime
from html.parser import HTMLParser
from typing import Any
from urllib.parse import urljoin, urlsplit
from zipfile import BadZipFile

from openpyxl import load_workbook

from macrolab.data.sources.http import SourceClient, validate_url

COMMODITY_PAGE = "https://www.worldbank.org/en/research/commodity-markets"
WDI_BASE = "https://api.worldbank.org/v2"
TERMS_STATUS = "unresolved – no redistribution of raw data"
STALE_MONTHS = 3
PROFILE_INDICATORS = {
    "SP.POP.TOTL": ("population", "persons"),
    "NY.GDP.PCAP.PP.KD": ("gdp_per_capita_ppp", "constant international dollars"),
    "NE.TRD.GNFS.ZS": ("trade_openness", "% of GDP"),
    "GB.XPD.RSDV.GD.ZS": ("rd_gdp", "% of GDP"),
    "SP.POP.DPND": ("age_dependency", "% of working-age population"),
}


def _issue(code: str, message: str, **details: Any) -> dict[str, Any]:
    return {"source": "WorldBank", "code": code, "message": message, **details}


def _as_of(client: SourceClient) -> date:
    value = getattr(client, "as_of_date", None)
    if value is None:
        return datetime.now(UTC).date()
    return date.fromisoformat(str(value)[:10])


def _finite(value: Any) -> bool:
    try:
        return value is not None and not isinstance(value, bool) and math.isfinite(float(value))
    except (ValueError, TypeError):
        return False


class _DownloadLinks(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self.links.extend(value for key, value in attrs if key == "href" and value)


def discover_monthly_url(body: bytes, page_url: str = COMMODITY_PAGE) -> str:
    """Resolve a unique current monthly workbook link; never fall back to an old URL."""
    parser = _DownloadLinks()
    parser.feed(body.decode("utf-8", errors="replace"))
    candidates = {
        urljoin(page_url, link)
        for link in parser.links
        if urlsplit(link).path.lower().endswith("cmo-historical-data-monthly.xlsx")
    }
    if len(candidates) != 1:
        raise ValueError(f"Expected one distinct monthly workbook link, found {len(candidates)}")
    result = candidates.pop()
    validate_url(result)
    if urlsplit(result).hostname not in {"thedocs.worldbank.org", "www.worldbank.org"}:
        raise ValueError("Monthly workbook link is not hosted by the World Bank")
    return result


def _normal_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("*", "")).strip().lower()


def parse_monthly_workbook(
    body: bytes, source_url: str, raw_path: str | None, as_of: date
) -> dict[str, Any]:
    """Read named columns, retaining each observation's original cell coordinates."""
    rows: list[dict[str, Any]] = []
    variants: list[dict[str, Any]] = []
    evidence: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        mappings = [
            ("Monthly Prices", "Crude oil, Brent", "oil_brent", None, "USD/bbl", None),
            ("Monthly Indices", "Non-energy", "commod_idx", "non_energy", "index", "2010"),
            ("Monthly Indices", "Energy", "commod_idx", "energy", "index", "2010"),
        ]
        for sheet_name, header, concept, variant, unit, base_year in mappings:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Required workbook sheet absent: {sheet_name}")
            values = list(workbook[sheet_name].iter_rows(values_only=True))
            cells = [
                (r, c)
                for r, row in enumerate(values[:20])
                for c, cell in enumerate(row)
                if _normal_header(cell) == _normal_header(header)
            ]
            if len(cells) != 1:
                raise ValueError(f"Ambiguous or missing workbook header: {sheet_name}/{header}")
            header_row, column = cells[0]
            if base_year and not any("2010=100" in str(row[0]) for row in values[:20]):
                raise ValueError(f"Index base year requires review: {sheet_name}/{header}")
            if concept == "oil_brent" and "$/bbl" not in str(values[header_row + 1][column]):
                raise ValueError("Brent unit no longer matches dollars per barrel")
            column_rows: list[dict[str, Any]] = []
            missing = 0
            for row_number, row in enumerate(values[header_row + 1 :], header_row + 2):
                source_period = str(row[0] or "")
                match = re.fullmatch(r"(\d{4})M(0[1-9]|1[0-2])", source_period)
                if not match:
                    continue
                value = row[column] if len(row) > column else None
                if not _finite(value):
                    missing += 1
                    continue
                period = f"{match[1]}-{match[2]}"
                column_rows.append(
                    {
                        "country": "GLOBAL",
                        "concept_id": concept,
                        "source": "WorldBank",
                        "source_series_id": f"PinkSheet/{sheet_name}/{header}",
                        "date": period,
                        "value": float(str(value)),
                        "unit": unit,
                        "unit_mult": 0,
                        "sa_flag": "UNKNOWN",
                        "price_basis": "nominal USD",
                        "base_year": base_year,
                        "freq": "M",
                        "obs_status": "reported",
                        "source_url": source_url,
                        "attribution": "World Bank, Commodity Price Data (The Pink Sheet)",
                        "terms_status": TERMS_STATUS,
                        "flags": ["seasonal_adjustment_not_reported"],
                        "variant": variant,
                        "raw_path": raw_path,
                        "audit": {
                            "format": "xlsx",
                            "sheet": sheet_name,
                            "column": column + 1,
                            "column_name": str(values[header_row][column]),
                            "row": row_number,
                            "source_period": source_period,
                        },
                    }
                )
            if not column_rows:
                raise ValueError(f"No finite monthly observations: {sheet_name}/{header}")
            periods = [item["date"] for item in column_rows]
            if len(periods) != len(set(periods)):
                raise ValueError(f"Duplicate monthly dates require review: {sheet_name}/{header}")
            latest = max(periods)
            age = (as_of.year - int(latest[:4])) * 12 + as_of.month - int(latest[5:7])
            stale = age > STALE_MONTHS
            for item in column_rows:
                if stale:
                    item["flags"].append("stale_source")
                if item["date"] > as_of.strftime("%Y-%m"):
                    item["flags"].append("future_observation_date")
            if stale:
                issues.append(
                    _issue(
                        "stale_pink_sheet",
                        f"Pink Sheet {header} ends {latest}, older than {STALE_MONTHS} months.",
                        concept_id=concept,
                        variant=variant,
                        latest_period=latest,
                    )
                )
            evidence.append(
                {
                    "concept_id": concept,
                    "variant": variant,
                    "sheet": sheet_name,
                    "header": str(values[header_row][column]),
                    "header_row": header_row + 1,
                    "column": column + 1,
                    "first_period": min(periods),
                    "latest_period": latest,
                    "observations": len(column_rows),
                    "missing_cells": missing,
                    "as_of_date": as_of.isoformat(),
                    "latest_age_months": age,
                    "stale_threshold_months": STALE_MONTHS,
                    "stale": stale,
                    "source_url": source_url,
                    "raw_path": raw_path,
                }
            )
            (variants if variant == "energy" else rows).extend(column_rows)
    finally:
        workbook.close()
    return {"rows": rows, "variant_rows": variants, "issues": issues, "mapping_evidence": evidence}


def parse_wdi_profiles(
    body: bytes, source_url: str, raw_path: str | None, allowed_countries: set[str]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    payload = json.loads(body)
    if not isinstance(payload, list) or len(payload) != 2 or not isinstance(payload[1], list):
        raise ValueError("WDI returned no metadata/observation pair")
    metadata = payload[0]
    profiles: list[dict[str, Any]] = []
    for index, item in enumerate(payload[1]):
        source_country = item.get("countryiso3code")
        country = "EA" if source_country == "EMU" else source_country
        indicator = item.get("indicator", {}).get("id")
        if country not in allowed_countries or indicator not in PROFILE_INDICATORS:
            continue
        if not _finite(item.get("value")) or not re.fullmatch(r"\d{4}", str(item.get("date"))):
            continue
        profile_id, unit = PROFILE_INDICATORS[indicator]
        title = item["indicator"]["value"]
        ppp_base = re.search(r"constant (\d{4}) international", title)
        profiles.append(
            {
                "country": country,
                "profile_id": profile_id,
                "source": "WorldBank",
                "source_series_id": indicator,
                "indicator_title": title,
                "date": str(item["date"]),
                "value": float(item["value"]),
                "unit": unit,
                "unit_mult": 0,
                "base_year": ppp_base[1] if ppp_base else None,
                "freq": "A",
                "obs_status": item.get("obs_status") or "reported",
                "source_url": source_url,
                "raw_path": raw_path,
                "attribution": (
                    "World Bank, World Development Indicators; underlying providers apply"
                ),
                "terms_status": TERMS_STATUS,
                "source_last_updated": metadata.get("lastupdated"),
                "flags": ["worldbank_euro_area_composition"] if country == "EA" else [],
                "audit": {
                    "format": "wdi-json",
                    "json_pointer": f"/1/{index}/value",
                    "source_period": item["date"],
                    "source_country": source_country,
                    "indicator": indicator,
                },
            }
        )
    return profiles, metadata


def fetch_worldbank(client: SourceClient, countries: list[dict[str, Any]]) -> dict[str, Any]:
    """Fetch current Pink Sheet and annual profiles, leaving partial failures visible."""
    result: dict[str, Any] = {
        "rows": [],
        "variant_rows": [],
        "profiles": [],
        "issues": [],
        "mapping_evidence": [],
    }
    page = client.fetch("WorldBank", "pink-current-page", COMMODITY_PAGE, "text/html")
    try:
        if page.record.status != 200 or page.body is None:
            raise ValueError(f"Commodity page unavailable: {page.record.error}")
        url = discover_monthly_url(page.body, page.record.final_url or COMMODITY_PAGE)
        workbook = client.fetch(
            "WorldBank",
            "pink-current-monthly",
            url,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        if workbook.record.status != 200 or workbook.body is None:
            raise ValueError(f"Current monthly workbook unavailable: {workbook.record.error}")
        parsed = parse_monthly_workbook(
            workbook.body,
            workbook.record.final_url or url,
            workbook.record.raw_path,
            _as_of(client),
        )
        for key in parsed:
            result[key].extend(parsed[key])
        result["download_resolution"] = {
            "page_url": COMMODITY_PAGE,
            "page_raw_path": page.record.raw_path,
            "resolved_url": url,
            "workbook_sha256": workbook.record.sha256,
            "as_of_date": _as_of(client).isoformat(),
        }
    except (ValueError, KeyError, BadZipFile, OSError) as exc:
        result["issues"].append(_issue("current_pink_sheet_unavailable", str(exc)))

    codes = sorted({str(country["iso3"]) for country in countries})
    source_codes = ["EMU" if code == "EA" else code for code in codes]
    if not source_codes:
        result["issues"].append(
            _issue("profile_country_list_empty", "No runtime countries supplied.")
        )
        return result
    indicator_codes = ";".join(PROFILE_INDICATORS)
    base_url = (
        f"{WDI_BASE}/country/{';'.join(source_codes)}/indicator/{indicator_codes}"
        f"?source=2&date=1999:{_as_of(client).year}&format=json&per_page=20000"
    )
    page_number = 1
    while True:
        url = f"{base_url}&page={page_number}"
        response = client.fetch(
            "WorldBank", f"wdi-profiles-p{page_number}", url, "application/json"
        )
        try:
            if response.record.status != 200 or response.body is None:
                raise ValueError(f"WDI page {page_number} unavailable: {response.record.error}")
            profiles, metadata = parse_wdi_profiles(
                response.body, url, response.record.raw_path, set(codes)
            )
            result["profiles"].extend(profiles)
            pages = int(metadata.get("pages", 1))
            if pages > 100:
                raise ValueError("WDI pagination exceeds bounded profile request")
            if page_number >= pages:
                break
            page_number += 1
        except (ValueError, KeyError, TypeError) as exc:
            result["issues"].append(_issue("wdi_profile_fetch_failed", str(exc)))
            break
    covered = {(item["country"], item["profile_id"]) for item in result["profiles"]}
    for country in codes:
        for profile_id, _unit in PROFILE_INDICATORS.values():
            if (country, profile_id) not in covered:
                result["issues"].append(
                    _issue(
                        "missing_profile", f"No WDI {profile_id} for {country}.", country=country
                    )
                )
    return result
