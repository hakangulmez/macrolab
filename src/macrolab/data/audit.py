"""Independently re-read official raw payloads for country/date spot checks."""

from __future__ import annotations

import csv
import io
import json
import math
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from macrolab.data.standardise import normalise_period


class SourceAudit:
    """Reference parsing is separate from the source adapters and their tidy output."""

    def __init__(self, root: Path) -> None:
        self.root = root
        self.cache: dict[str, Any] = {}

    def table(self, raw_path: str, parser: str) -> Any:
        if raw_path in self.cache:
            return self.cache[raw_path]
        path = (self.root / raw_path).resolve()
        if not path.is_relative_to((self.root / "data/raw").resolve()):
            raise ValueError("Audit reference must be an archived official response")
        body = path.read_bytes()
        if parser == "csv":
            result: Any = list(csv.DictReader(io.StringIO(body.decode("utf-8-sig"))))
        elif parser in {"sdmx_xml", "sdmx-xml"}:
            tree = ET.fromstring(body)
            result = []
            for dataset in tree.iter():
                if dataset.tag.split("}")[-1] != "DataSet":
                    continue
                groups = [
                    element.attrib for element in dataset if element.tag.split("}")[-1] == "Group"
                ]
                for series in dataset:
                    if series.tag.split("}")[-1] != "Series":
                        continue
                    metadata = {
                        key: value
                        for key, value in dataset.attrib.items()
                        if not key.startswith("{")
                    }
                    for group in groups:
                        common = set(group) & set(series.attrib)
                        if common and all(group[key] == series.attrib[key] for key in common):
                            metadata.update(group)
                    metadata.update(series.attrib)
                    for child in series:
                        if child.tag.split("}")[-1] in {"SeriesKey", "Attributes"}:
                            metadata.update(
                                {
                                    item.attrib["id"]: item.attrib["value"]
                                    for item in child
                                    if "id" in item.attrib and "value" in item.attrib
                                }
                            )
                    for observation in series:
                        if observation.tag.split("}")[-1] != "Obs":
                            continue
                        value = {**metadata, **observation.attrib}
                        for item in observation:
                            tag = item.tag.split("}")[-1]
                            if tag == "ObsDimension":
                                value["TIME_PERIOD"] = item.attrib["value"]
                            elif tag == "ObsValue":
                                value["OBS_VALUE"] = item.attrib["value"]
                        result.append(value)
        elif parser in {"jsonstat", "json-stat", "wdi-json", "json"}:
            result = json.loads(body)
        elif parser == "xls":
            import pandas as pd

            frame = pd.read_excel(io.BytesIO(body), sheet_name="Data 1", header=None, engine="xlrd")
            result = frame.values.tolist()
        elif parser == "xlsx":
            workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            try:
                result = {
                    sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook
                }
            finally:
                workbook.close()
        else:
            raise ValueError(f"Unsupported independent audit parser: {parser}")
        self.cache[raw_path] = result
        return result

    def reference_value(self, row: dict[str, Any]) -> float:
        audit = row.get("audit") or {}
        parser = audit.get("parser", audit.get("format"))
        if parser == "derived_term_spread":
            long, short = audit["components"]
            return self.reference_value(long) - self.reference_value(short)
        if not row.get("raw_path"):
            raise ValueError("No raw response provenance for this observation")
        if not isinstance(parser, str):
            raise ValueError("Missing independent audit parser")
        table = self.table(row["raw_path"], parser)
        source_period = str(audit["source_period"])
        comparable_period = source_period.split("T")[0] if parser == "xls" else source_period
        if normalise_period(comparable_period, row["freq"]) != row["date"]:
            raise ValueError("Source period does not match standardised date")
        if parser in {"csv", "sdmx_xml", "sdmx-xml"}:
            time_key, value_key = (
                audit.get("time_column", "TIME_PERIOD"),
                audit.get("value_column", "OBS_VALUE"),
            )
            selector = audit["selector"]
            matches = [
                candidate
                for candidate in table
                if str(candidate.get(time_key)) == source_period
                and all(
                    str(candidate.get(key, "")) == str(value) for key, value in selector.items()
                )
            ]
            if not matches:
                # Some source adapters normalise IMF month labels while the raw file retains M.
                matches = [
                    candidate
                    for candidate in table
                    if candidate.get(time_key)
                    and all(
                        str(candidate.get(key, "")) == str(value) for key, value in selector.items()
                    )
                    and normalise_period(str(candidate[time_key]), row["freq"]) == row["date"]
                ]
            if len(matches) != 1:
                raise ValueError(f"Independent source selector matched {len(matches)} observations")
            return float(matches[0][value_key])
        if parser in {"jsonstat", "json-stat"}:
            index = int(audit["flat_index"])
            coordinates: dict[str, str] = {}
            dimension_names = table["id"]
            sizes = table["size"]
            if index < 0 or index >= math.prod(sizes):
                raise ValueError("JSON-stat audit index is out of bounds")
            for position, name in enumerate(dimension_names):
                coordinate = index // math.prod(sizes[position + 1 :]) % sizes[position]
                categories = table["dimension"][name]["category"]["index"]
                if isinstance(categories, list):
                    coordinates[name] = categories[coordinate]
                else:
                    coordinates[name] = next(
                        code for code, offset in categories.items() if offset == coordinate
                    )
            if coordinates.get("time") != source_period:
                raise ValueError("JSON-stat raw time coordinate does not match the audit date")
            if any(
                coordinates.get(key) != str(value)
                for key, value in audit.get("selector", {}).items()
            ):
                raise ValueError(
                    "JSON-stat raw dimension coordinates differ from the selected series"
                )
            values = table["value"]
            return float(values[str(index)] if isinstance(values, dict) else values[int(index)])
        if parser == "xls":
            import re

            cell = re.fullmatch(r"B(\d+)", audit["value_cell"])
            if not cell or audit["source_key"] != "RBRTE" or table[1][1] != "RBRTE":
                raise ValueError("Invalid EIA raw cell or source key")
            item = table[int(cell[1]) - 1]
            if item[0].strftime("%Y-%m") != row["date"]:
                raise ValueError("EIA raw cell date mismatch")
            return float(item[1])
        if parser == "xlsx":
            sheet = table[audit["sheet"]]
            source_row = sheet[int(audit["row"]) - 1]
            if str(source_row[0]) != source_period:
                raise ValueError("Workbook row date differs from audit reference")
            return float(source_row[int(audit["column"]) - 1])
        if parser == "wdi-json":
            keys = audit["json_pointer"].strip("/").split("/")
            if len(keys) != 3 or keys[0] != "1" or keys[-1] != "value":
                raise ValueError("WDI audit pointer must identify an observation value")
            observation = table[1][int(keys[1])]
            if str(observation["date"]) != source_period:
                raise ValueError("WDI raw date does not match the audit period")
            if observation.get("countryiso3code") != audit.get("source_country"):
                raise ValueError("WDI raw country does not match the audit country")
            if observation["indicator"]["id"] != audit.get("indicator"):
                raise ValueError("WDI raw indicator does not match the audit concept")
            return float(observation["value"])
        raise ValueError("Audit parser has no value accessor")


def audit_data(
    root: Path,
    rows: list[dict[str, Any]],
    concepts: list[dict[str, Any]],
    preferred_countries: dict[str, list[str]] | None = None,
) -> dict[str, Any]:
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[(row["concept_id"], row["country"])].append(row)
    reader = SourceAudit(root)
    results, concept_results = [], []
    for definition in concepts:
        concept = definition["concept_id"]
        eligible = sorted(
            country
            for (cid, country), values in groups.items()
            if cid == concept and len(values) >= 3 and country != "EA"
        )
        preferred = (preferred_countries or {}).get(concept, ["USA", "DEU", "JPN"])
        selected = [country for country in preferred if country in eligible][:3]
        selected += [country for country in eligible if country not in selected][
            : 3 - len(selected)
        ]
        for country in selected:
            values = sorted(groups[(concept, country)], key=lambda row: row["date"])
            for index in (0, len(values) // 2, len(values) - 1):
                row = values[index]
                record = {
                    "country": country,
                    "concept_id": concept,
                    "date": row["date"],
                    "source": row["source"],
                    "source_url": row["source_url"],
                    "source_series_id": row["source_series_id"],
                    "raw_path": row.get("raw_path"),
                    "status": "unavailable",
                    "error": None,
                    "absolute_difference": None,
                    "global_factor_shared": definition["scope"] == "global",
                }
                try:
                    reference = reader.reference_value(row)
                    difference = abs(row["value"] - reference)
                    record.update(
                        {
                            "status": "matched"
                            if math.isclose(row["value"], reference, rel_tol=1e-10, abs_tol=1e-8)
                            else "mismatch",
                            "absolute_difference": difference,
                        }
                    )
                except (ValueError, TypeError, KeyError, IndexError, OSError) as exc:
                    record["error"] = str(exc)
                results.append(record)
        relevant = [record for record in results if record["concept_id"] == concept]
        concept_results.append(
            {
                "concept_id": concept,
                "countries": selected,
                "required_checks": 9,
                "completed_checks": len(relevant),
                "matched": sum(record["status"] == "matched" for record in relevant),
                "mismatches": sum(record["status"] == "mismatch" for record in relevant),
                "unavailable": sum(record["status"] == "unavailable" for record in relevant),
                "status": "passed"
                if len(relevant) == 9 and all(record["status"] == "matched" for record in relevant)
                else "incomplete",
            }
        )
    return {
        "method": (
            "Independent re-read of archived official APIs/files; first, middle and last "
            "observed date for three countries; the listed per-concept countries are "
            "preferred (default USA/DEU/JPN), else alphabetical available countries."
        ),
        "global_factor_note": (
            "Global factors are identical shared inputs; three country checks do not "
            "represent three independent source measurements."
        ),
        "preferred_countries": preferred_countries or {},
        "tolerance_absolute": 1e-8,
        "tolerance_relative": 1e-10,
        "required_checks": 9 * len(concepts),
        "performed_checks": len(results),
        "matched": sum(record["status"] == "matched" for record in results),
        "mismatches": sum(record["status"] == "mismatch" for record in results),
        "unavailable": sum(record["status"] == "unavailable" for record in results),
        "passed": all(record["status"] == "passed" for record in concept_results),
        "concepts": concept_results,
        "checks": results,
    }
