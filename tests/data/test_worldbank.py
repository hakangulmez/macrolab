"""Offline, synthetic World Bank fixtures: no provider observations are committed."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import pytest
from openpyxl import Workbook, load_workbook

from macrolab.data.countries import enrich_country_metadata, parse_country_metadata
from macrolab.data.sources.http import SourceClient
from macrolab.data.sources.worldbank import (
    discover_monthly_url,
    fetch_worldbank,
    parse_monthly_workbook,
    parse_wdi_profiles,
)

URL = "https://thedocs.worldbank.org/en/doc/synthetic/CMO-Historical-Data-Monthly.xlsx"


def workbook_bytes(*, last_period: str = "2026M09", brent_unit: str = "($/bbl)") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Monthly Prices"
    sheet.append(["Synthetic data"])
    sheet.append([None, "Other commodity", "Crude oil, Brent"])
    sheet.append([None, "($/kg)", brent_unit])
    sheet.append(["2026M01", 1, 10])
    sheet.append(["2026M02", 2, "…"])
    sheet.append([last_period, 3, 20])
    indices = workbook.create_sheet("Monthly Indices")
    indices.append(["monthly indices based on nominal US dollars, 2010=100"])
    indices.append([None, "Non-energy **", "Energy"])
    indices.append(["2026M01", 100, 200])
    indices.append([last_period, 105, 210])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_current_link_resolution_requires_unique_official_monthly_file() -> None:
    page = f'<a href="{URL}">Monthly</a><a href="{URL}">Repeated navigation</a>'.encode()
    assert discover_monthly_url(page) == URL
    with pytest.raises(ValueError, match="found 0"):
        discover_monthly_url(b'<a href="old-annual.xlsx">Annual</a>')
    with pytest.raises(ValueError, match="found 2"):
        discover_monthly_url(page + page.replace(b"/synthetic/", b"/ambiguous/"))
    with pytest.raises(ValueError):
        discover_monthly_url(page.replace(b"thedocs.worldbank.org", b"unapproved.example"))


def test_workbook_preserves_variants_missing_cells_and_auditable_coordinates() -> None:
    body = workbook_bytes()
    parsed = parse_monthly_workbook(body, URL, "raw/example.bin", date(2026, 10, 20))
    assert len(parsed["rows"]) == 4
    assert len(parsed["variant_rows"]) == 2
    assert {r["variant"] for r in parsed["variant_rows"]} == {"energy"}
    assert not any(r["date"] == "2026-02" for r in parsed["rows"])
    assert not any("stale_source" in r["flags"] for r in parsed["rows"])
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        for row in parsed["rows"] + parsed["variant_rows"]:
            audit = row["audit"]
            assert (
                workbook[audit["sheet"]].cell(audit["row"], audit["column"]).value == row["value"]
            )
    finally:
        workbook.close()
    assert parsed["mapping_evidence"][0]["missing_cells"] == 1


@pytest.mark.parametrize(("month", "stale"), [("2026M07", False), ("2026M06", True)])
def test_staleness_threshold_is_three_elapsed_calendar_months(month: str, stale: bool) -> None:
    parsed = parse_monthly_workbook(
        workbook_bytes(last_period=month), URL, None, date(2026, 10, 20)
    )
    assert all(item["stale"] is stale for item in parsed["mapping_evidence"])
    assert all(("stale_source" in row["flags"]) is stale for row in parsed["rows"])
    assert bool(parsed["issues"]) is stale


def test_changed_unit_and_duplicate_periods_fail_instead_of_silent_remapping() -> None:
    with pytest.raises(ValueError, match="Brent unit"):
        parse_monthly_workbook(workbook_bytes(brent_unit="EUR/bbl"), URL, None, date(2026, 10, 20))
    with pytest.raises(ValueError, match="Duplicate"):
        parse_monthly_workbook(workbook_bytes(last_period="2026M01"), URL, None, date(2026, 10, 20))


def test_wdi_nulls_stay_missing_and_profile_unit_and_audit_survive() -> None:
    items = [
        {
            "countryiso3code": "EMU",
            "indicator": {
                "id": "NY.GDP.PCAP.PP.KD",
                "value": "GDP per capita, PPP (constant 2021 international $)",
            },
            "date": "2025",
            "value": 123,
        },
        {
            "countryiso3code": "EMU",
            "indicator": {"id": "SP.POP.TOTL", "value": "Population, total"},
            "date": "2024",
            "value": None,
        },
        {
            "countryiso3code": "USA",
            "indicator": {"id": "SP.POP.TOTL", "value": "Population, total"},
            "date": "2025",
            "value": 456,
        },
    ]
    body = json.dumps([{"pages": 1, "lastupdated": "2026-10-08"}, items]).encode()
    profiles, _metadata = parse_wdi_profiles(
        body, "https://api.worldbank.org/example", None, {"EA"}
    )
    assert len(profiles) == 1
    assert profiles[0]["country"] == "EA"
    assert profiles[0]["base_year"] == "2021"
    assert profiles[0]["date"] == "2025"
    assert profiles[0]["audit"]["json_pointer"] == "/1/0/value"


def rich_body() -> bytes:
    return json.dumps(
        {
            "pages": 1,
            "source": [
                {
                    "concept": [
                        {
                            "id": "Country",
                            "variable": [
                                {
                                    "id": "GRC",
                                    "metatype": [
                                        {"id": "2-alphacode", "value": "GR"},
                                        {"id": "CurrencyUnit", "value": "Euro"},
                                        {"id": "Region", "value": "WB region"},
                                    ],
                                },
                                {
                                    "id": "TUR",
                                    "metatype": [
                                        {"id": "2-alphacode", "value": "TR"},
                                        {"id": "CurrencyUnit", "value": "New Turkish lira"},
                                    ],
                                },
                                {
                                    "id": "EMU",
                                    "metatype": [
                                        {"id": "CurrencyUnit", "value": "Euro"},
                                    ],
                                },
                            ],
                        }
                    ]
                }
            ],
        }
    ).encode()


class MetadataClient:
    def __init__(self, root: Path) -> None:
        self.root = root
        self.as_of_date = "2026-10-20"

    def fetch(self, _provider: str, request_id: str, _url: str, _accept: str) -> Any:
        return SimpleNamespace(
            body=rich_body() if request_id == "country-rich-metadata" else b"Official ECB page",
            record=SimpleNamespace(status=200, error=None, raw_path="raw/synthetic.bin"),
        )


def test_country_enrichment_preserves_runtime_roster_and_uses_adoption_sample() -> None:
    root = Path(__file__).resolve().parents[2]
    client = cast(SourceClient, MetadataClient(root))
    originals = [
        {"iso3": "GRC", "name": "Greece", "oecd_region": "OECD Europe"},
        {"iso3": "TUR", "name": "Türkiye"},
    ]
    enriched = enrich_country_metadata(originals, client)
    assert [row["iso3"] for row in enriched] == ["GRC", "TUR"]
    assert "currency" not in originals[0]
    assert enriched[0]["default_shock_sample_start"] == "2001-01-01"
    assert enriched[0]["region"] == "OECD Europe"
    assert enriched[0]["worldbank_region"] == "WB region"
    assert enriched[1]["currency"] == "Turkish lira"
    assert enriched[1]["monetary_union"] is None
    assert all(note["status"] == "DRAFT" for row in enriched for note in row["notes"])
    assert parse_country_metadata(rich_body())["GRC"]["2-alphacode"] == "GR"


def test_euro_area_enrichment_preserves_eu_methodology_eligibility() -> None:
    root = Path(__file__).resolve().parents[2]
    original = {"iso3": "EA", "name": "Euro area", "eu_member": True, "group": "aggregate"}
    enriched = enrich_country_metadata([original], cast(SourceClient, MetadataClient(root)))[0]
    assert enriched["eu_member"] is True
    assert enriched["monetary_union"] == "EA"
    assert enriched["currency_code"] == "EUR"
    assert enriched["iso2_status"] == "internal_aggregate_code"


def test_missing_current_download_produces_issue_without_stale_fallback() -> None:
    class UnavailableClient:
        as_of_date = "2026-10-20"

        def fetch(self, *_args: Any) -> Any:
            return SimpleNamespace(body=None, record=SimpleNamespace(status=503, error="HTTP 503"))

    result = fetch_worldbank(cast(SourceClient, UnavailableClient()), [])
    assert result["rows"] == []
    assert result["variant_rows"] == []
    assert result["issues"][0]["code"] == "current_pink_sheet_unavailable"
